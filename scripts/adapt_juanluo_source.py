#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
adapt_juanluo_source.py — 把《全样本生产数据》长表 适配成 老宽表，供 update_juanluo_source.py 消费

背景：
  Mysteel 有**两种**导出格式：
    A) 老宽表（一个文件两 sheet：螺纹全样本 / 热卷全样本；col1=年 col2=周 col3=日期，col4+ 区域块）
       —— update_juanluo_source.py 直接吃。
    B) 新长表（**两个文件**：《螺纹钢全样本生产数据YYYY-M-D.xlsx》《热卷全样本生产数据YYYY-M-D.xlsx》；
       sheet「…分区域数据」，列 = 数据日期|品种|区域|开工率|产能利用率|周产量|钢厂库存；一次给最近 4 周）
       —— 本脚本把它转成 A 的临时宽表。

转换规则：
  · 区域顺序统一为 东北,华北,华东,华南,华中,西北,西南,合计（"总计"/"全国" → "合计"；去掉"区"后缀）
  · 输出 col1=年, col2=周(ISO周), col3=日期(datetime), col4..35 = 8 块 × 4 指标
  · 前两行为表头（区域名 / 指标名），数据自第 3 行起（与老宽表一致，update 脚本从第 3 行扫日期）

用法：
  python adapt_juanluo_source.py --out C:/tmp/wide.xlsx --luowen "<螺纹长表.xlsx>" --hot "<热卷长表.xlsx>"
  python adapt_juanluo_source.py --out C:/tmp/wide.xlsx --discover      # 自动找微信最新两份
"""
from __future__ import annotations

import argparse
import datetime
import glob
import os
import re

import openpyxl
from openpyxl.styles import Alignment

REGIONS = ["东北", "华北", "华东", "华南", "华中", "西北", "西南", "合计"]
ALIAS = {"总计": "合计", "全国": "合计", "合计": "合计"}
METRICS = ["开工率", "产能利用率", "周产量", "钢厂库存"]


def norm_region(raw):
    if raw is None:
        return None
    s = str(raw).strip()
    s = re.sub(r"[（(].*?[)）]", "", s).strip()
    if s in ALIAS:
        return ALIAS[s]
    s = s.rstrip("区").strip()
    return s if s in REGIONS else None


def read_long(path):
    """读长表 → {date: {region: [4 指标]}}，并返回品种名"""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    out, variety = {}, None
    try:
        for sn in wb.sheetnames:
            ws = wb[sn]
            rows = list(ws.iter_rows(values_only=True))
            if not rows:
                continue
            hdr = [str(c).strip() if c is not None else "" for c in rows[0]]
            if not any("数据日期" in h for h in hdr) or not any("区域" in h for h in hdr):
                continue
            ci = {name: i for i, name in enumerate(hdr)}
            # 定位各列（列名可能带单位后缀，如「周产量（万吨）」）
            def col_of(keyword, exclude=None):
                for name, i in ci.items():
                    if keyword in name and (exclude is None or exclude not in name):
                        return i
                return None

            c_date = col_of("数据日期")
            c_var = col_of("品种")
            c_reg = col_of("区域")
            c_open = col_of("开工率")
            c_cap = col_of("产能利用率")
            c_prod = col_of("周产量")
            c_stock = col_of("钢厂库存")

            for r in rows[1:]:
                cells = list(r)
                if all(c is None for c in cells):
                    continue
                d = cells[c_date] if c_date is not None and c_date < len(cells) else None
                if isinstance(d, datetime.datetime):
                    d = d.date()
                elif isinstance(d, datetime.date):
                    pass
                else:
                    continue
                reg = norm_region(cells[c_reg] if c_reg is not None and c_reg < len(cells) else None)
                if reg is None:
                    continue
                if variety is None and c_var is not None and c_var < len(cells):
                    v = cells[c_var]
                    if v:
                        variety = str(v).strip()
                vals = []
                for ci2 in (c_open, c_cap, c_prod, c_stock):
                    v = cells[ci2] if ci2 is not None and ci2 < len(cells) else None
                    vals.append(float(v) if isinstance(v, (int, float)) else None)
                out.setdefault(d, {})[reg] = vals
    finally:
        wb.close()

    # 品种兜底：从文件名猜
    if not variety:
        base = os.path.basename(path)
        variety = "螺纹钢" if "螺纹" in base else ("热轧板卷" if "热卷" in base else None)
    return out, variety


def guess_kind(path, variety=None):
    key = (variety or "") + " " + os.path.basename(path)
    if "螺纹" in key:
        return "luowen"
    if "热卷" in key or "热轧" in key:
        return "hot"
    return None


def is_legacy(path):
    """老宽表：含 sheet 名「螺纹全样本」「热卷全样本」之一"""
    try:
        wb = openpyxl.load_workbook(path, read_only=True)
        names = wb.sheetnames
        wb.close()
        return ("螺纹全样本" in names) or ("热卷全样本" in names)
    except Exception:
        return False


def discover_latest():
    """在微信落点里找最新的 螺纹/热卷 长表，返回 (luowen_path, hot_path)"""
    roots = [
        r"D:\微信\xwechat_files\*\msg\file\*\*全样本生产数据*.xlsx",
        r"D:\微信\xwechat_files\*\temp\RWTemp\*\*\*全样本生产数据*.xlsx",
    ]
    cands = []
    for pat in roots:
        cands += glob.glob(pat)
    lw = [p for p in cands if "螺纹" in os.path.basename(p)]
    ht = [p for p in cands if "热卷" in os.path.basename(p)]
    lw.sort(key=os.path.getmtime)
    ht.sort(key=os.path.getmtime)
    return (lw[-1] if lw else None), (ht[-1] if ht else None)


def write_wide(datasets, out_path):
    """datasets: {'luowen': {date:{reg:[...]}}, 'hot': {...}} → 写老宽表形态 xlsx"""
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    sheet_name = {"luowen": "螺纹全样本", "hot": "热卷全样本"}
    total_rows = 0
    for kind in ("luowen", "hot"):
        data = datasets.get(kind)
        if not data:
            continue
        ws = wb.create_sheet(sheet_name[kind])
        # 行1：区域名（每块 4 列，首列写区域名）
        ws.cell(1, 1, "")
        ws.cell(1, 2, "")
        ws.cell(1, 3, "")
        for i, reg in enumerate(REGIONS):
            ws.cell(1, 4 + 4 * i, reg)
        # 行2：指标名
        ws.cell(2, 1, "年")
        ws.cell(2, 2, "周")
        ws.cell(2, 3, "日期")
        for i in range(len(REGIONS)):
            for j, m in enumerate(METRICS):
                ws.cell(2, 4 + 4 * i + j, m)
        # 数据
        r = 3
        for d in sorted(data.keys()):
            regions = data[d]
            if not regions:
                continue
            ws.cell(r, 1, d.year)
            ws.cell(r, 2, d.isocalendar()[1])
            c = ws.cell(r, 3, datetime.datetime(d.year, d.month, d.day))
            c.number_format = "yyyy-mm-dd"
            for i, reg in enumerate(REGIONS):
                vals = regions.get(reg)
                if not vals:
                    continue
                for j, v in enumerate(vals):
                    ws.cell(r, 4 + 4 * i + j, v)
            r += 1
        total_rows += (r - 3)
        ws.freeze_panes = "D3"
        ws.column_dimensions["A"].width = 6
        ws.column_dimensions["B"].width = 5
        ws.column_dimensions["C"].width = 12
    if not wb.sheetnames:
        raise SystemExit("❌ 没能解析出任何数据（两份长表都为空的？）")
    os.makedirs(os.path.dirname(os.path.abspath(out_path)), exist_ok=True)
    wb.save(out_path)
    return total_rows


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--out", required=True, help="输出的临时宽表 xlsx 路径")
    ap.add_argument("--luowen", help="螺纹钢全样本生产数据*.xlsx")
    ap.add_argument("--hot", help="热卷全样本生产数据*.xlsx")
    ap.add_argument("--discover", action="store_true", help="自动在微信落点找最新两份长表")
    args = ap.parse_args()

    lw, ht = args.luowen, args.hot
    if args.discover or (not lw and not ht):
        dl, dh = discover_latest()
        lw = lw or dl
        ht = ht or dh
        print(f"[discover] 螺纹={os.path.basename(lw) if lw else None}")
        print(f"[discover] 热卷={os.path.basename(ht) if ht else None}")

    if not lw and not ht:
        raise SystemExit("❌ 未提供也未发现任何长表源文件")

    # 若给的就是老宽表 → 原样复制
    if lw and ht is None and is_legacy(lw):
        import shutil
        os.makedirs(os.path.dirname(os.path.abspath(args.out)), exist_ok=True)
        shutil.copy2(lw, args.out)
        print(f"[legacy] 直接使用老宽表: {os.path.basename(lw)}")
        return

    datasets = {}
    for kind, path in (("luowen", lw), ("hot", ht)):
        if not path:
            continue
        if not os.path.exists(path):
            raise SystemExit(f"❌ 文件不存在: {path}")
        data, variety = read_long(path)
        k = guess_kind(path, variety) or kind
        if k in datasets:
            raise SystemExit(f"❌ 品种冲突（{k} 出现两次）: {path}")
        datasets[k] = data
        ds = sorted(data.keys())
        print(f"[read] {os.path.basename(path)} → {k} / {variety} / {len(ds)} 周 "
              f"({ds[0] if ds else '-'} ~ {ds[-1] if ds else '-'})")

    n = write_wide(datasets, args.out)
    print(f"[ok] 宽表已生成: {args.out}（数据行合计 {n}）")


if __name__ == "__main__":
    main()
